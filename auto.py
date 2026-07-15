import subprocess
import json
import re
import sys
from pathlib import Path

INPUT_DIR = Path("input")
SUMMARY_FILE = Path("/mnt/c/Users/Ferna/Desktop/database/SiPM_Data_Tools/checked/summary.xlsx")

# =========================
# LOAD / SAVE STATE (JSON)
# =========================
def load_state(state_file):
    if state_file.exists():
        return json.loads(state_file.read_text())
    return {}

def save_state(state, state_file):
    state_file.write_text(json.dumps(state, indent=2))

# =========================
# RUN COMMANDS
# =========================
def run(cmd):
    print(f"\n▶ {' '.join(cmd)}\n")
    subprocess.run(cmd, check=True)

# =========================
# UPDATE SUMMARY.XLSX
# =========================
def update_summary(box_id, tray_name, status="upload"):
    try:
        import openpyxl
    except ImportError:
        print("[summary] WARN: openpyxl no instalado, salto update")
        return

    if not SUMMARY_FILE.exists():
        print(f"[summary] WARN: {SUMMARY_FILE} no existe, salto update")
        return

    m = re.match(r"Box(\d+)", box_id)
    m2 = re.match(r"Tray0*(\d+)", tray_name, re.IGNORECASE)
    if not (m and m2):
        print(f"[summary] WARN: no puedo parsear {box_id} / {tray_name}")
        return
    box_num = int(m.group(1))
    tray_num = int(m2.group(1))

    try:
        wb = openpyxl.load_workbook(SUMMARY_FILE)
    except PermissionError:
        print(f"[summary] WARN: summary.xlsx bloqueado, cierralo en Excel")
        return

    ws = wb.active
    header = {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}
    col_box = header.get("Box")
    col_tray = header.get("Tray")
    col_status = header.get("Status")
    if not (col_box and col_tray and col_status):
        print("[summary] WARN: faltan columnas Box/Tray/Status")
        wb.close()
        return

    updated = False
    upload_fill = openpyxl.styles.PatternFill(
        patternType='solid',
        fgColor=openpyxl.styles.Color(theme=3, tint=0.3999755851924192),
    )
    for row in ws.iter_rows(min_row=2):
        cell_box = row[col_box - 1].value
        cell_tray = row[col_tray - 1].value
        if cell_box == box_num and cell_tray == tray_num:
            row[col_status - 1].value = status
            row[col_status - 1].fill = upload_fill
            updated = True
            break

    if updated:
        import tempfile, os, shutil
        try:
            with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
                tmp_path = tmp.name
            wb.save(tmp_path)
            shutil.copyfile(tmp_path, SUMMARY_FILE)
            os.unlink(tmp_path)
            print(f"[summary] {box_id}/{tray_name} -> {status}")
        except PermissionError as e:
            print(f"[summary] WARN PermissionError: {e}")
        except Exception as e:
            print(f"[summary] WARN {type(e).__name__}: {e}")
    else:
        print(f"[summary] WARN: fila {box_id}/{tray_name} no encontrada")
    wb.close()

# =========================
# UPDATE DOCKET
# =========================
def update_docket(tray_rel, docket_file):
    with open(docket_file, "r") as f:
        content = f.read()

    content = re.sub(
        r'box\s*=\s*".*?"',
        f'box = "{tray_rel}"',
        content
    )

    with open(docket_file, "w") as f:
        f.write(content)

# =========================
# MAIN PIPELINE
# =========================
def main():
    # preguntar a que base conectar
    while True:
        choice = input("A que base conectarse? (dev/prod): ").strip().lower()
        if choice in ("dev", "prod"):
            if choice == "prod":
                confirm = input("⚠️  ATENCION: vas a subir a PRODUCCION. Confirmas? (si/no): ").strip().lower()
                if confirm not in ("si", "s", "yes", "y"):
                    print("Cancelado. Vuelve a elegir.")
                    continue
            break
        print("Opcion invalida. Escribe 'dev' o 'prod'.")

    docket_file = f"docket_{choice}.py"
    state_file = Path(f"done_{choice}.json")
    print(f"\nUsando docket: {docket_file}  |  estado: {state_file}")

    state = load_state(state_file)

    if not INPUT_DIR.exists():
        print(f"ERROR: no existe directorio {INPUT_DIR}")
        return

    # detectar boxes en input/ agrupando por prefix (Box08, Box11, etc.)
    dirs = sorted([d for d in INPUT_DIR.iterdir() if d.is_dir()])

    boxes = {}
    for d in dirs:
        name = d.name
        m = re.match(r"(Box\d+)", name)
        if not m:
            continue
        box_id = m.group(1)
        if box_id not in boxes:
            boxes[box_id] = []
        boxes[box_id].append(d)

    if not boxes:
        print("No se han detectado cajas en input/")
        return

    # modo directo: python auto.py Box08 Tray000097
    if len(sys.argv) >= 3:
        target_box = sys.argv[1]
        target_tray = sys.argv[2]

        # buscar box por prefijo (Box08 => Box08_upload, Box08-checked, etc.)
        box_id = None
        for bid in boxes:
            if bid.lower() == target_box.lower():
                box_id = bid
                break
        if box_id is None:
            print(f"Box '{target_box}' no encontrado en input/")
            print(f"Boxes disponibles: {list(boxes.keys())}")
            return

        box_dirs = boxes[box_id]

        # buscar la tray por prefijo (Tray000097 => Tray000097_upload, etc.)
        found = None
        for box_dir in box_dirs:
            for d in box_dir.iterdir():
                if d.is_dir() and d.name.lower().startswith(target_tray.lower()):
                    found = (box_dir, d)
                    break
            if found:
                break

        if found is None:
            print(f"Tray '{target_tray}' no encontrada en {box_id}")
            return

        box_dir, tray_dir = found
        tray_name = tray_dir.name
        tray_rel = f"{box_dir.name}/{tray_name}"

        print(f"\nSubida directa: {tray_rel}")

        if tray_name in state.get(box_id, []):
            print(f"✔ SKIP {tray_rel} (ya esta subida)")
            return

        try:
            update_docket(tray_rel, docket_file)
            run(["hwdb-upload", docket_file])
            run(["hwdb-upload", docket_file, "--submit"])

            if box_id not in state:
                state[box_id] = []
            if tray_name not in state[box_id]:
                state[box_id].append(tray_name)
            save_state(state, state_file)
            update_summary(box_id, tray_name)

            print(f"✔ DONE {tray_rel}")

        except subprocess.CalledProcessError:
            print(f"\n❌ ERROR en {tray_rel}\n")
            return

    else:
        # menu interactivo para elegir caja
        box_ids = sorted(boxes.keys())
        print("\nCajas disponibles:")
        for i, bid in enumerate(box_ids, 1):
            print(f"  {i}) {bid}")

        while True:
            sel = input(f"\nSelecciona caja (1-{len(box_ids)}): ").strip()
            try:
                idx = int(sel) - 1
                if 0 <= idx < len(box_ids):
                    break
            except ValueError:
                pass
            print("Opcion invalida.")

        box_id = box_ids[idx]
        box_dirs = boxes[box_id]

        print(f"\nCaja seleccionada: {box_id}")

        print("\n==============================")
        print("HWDB BATCH START")
        print("==============================\n")

        done_trays = state.get(box_id, [])

        for box_dir in box_dirs:
            # solo directorios que empiecen por "Tray", ignora "results" y otros
            all_dirs = sorted([t for t in box_dir.iterdir() if t.is_dir()])
            trays = [t for t in all_dirs if t.name.lower().startswith("tray")]

            for tray in trays:
                tray_name = tray.name
                tray_rel = f"{box_dir.name}/{tray_name}"

                if tray_name in done_trays:
                    print(f"✔ SKIP {tray_rel}")
                    continue

                print(f"\nPROCESSING {tray_rel}")

                try:
                    update_docket(tray_rel, docket_file)
                    run(["hwdb-upload", docket_file])
                    run(["hwdb-upload", docket_file, "--submit"])

                    if box_id not in state:
                        state[box_id] = []
                    state[box_id].append(tray_name)
                    save_state(state, state_file)
                    update_summary(box_id, tray_name)

                    print(f"✔ DONE {tray_rel}")

                except subprocess.CalledProcessError:
                    print(f"\n❌ ERROR en {tray_rel} -> STOP\n")
                    return

        print("\n==============================")
        print("FINISHED")
        print("==============================\n")

if __name__ == "__main__":
    main()
